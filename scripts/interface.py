#!/app/jarvis-dev/venv/bin/python3
import json
import logging
import mimetypes
import socket
import re
import zipfile
from pathlib import Path
from datetime import datetime
from xml.etree import ElementTree as ET

from flask import Flask, request, jsonify, send_from_directory, make_response
from flask_sock import Sock
from werkzeug.utils import secure_filename

from config import (
    HOST,
    PORT,
    WEB_DIR,
    DATA_DIR,
    UPLOADS_DIR,
    DEFAULT_DISPLAY_NAME,
    render_greeting,
    ensure_directories,
    BRAIN_SOCKET,
    EARS_SOCKET,
    EYES_SOCKET,
    MOUTH_SOCKET,
    ADMIN_EMAILS,
)
from auth import AuthManager
from memory import MemoryStore
from location_service import validate_home_location

ensure_directories()

LOG_FILE = DATA_DIR / "logs" / "interface.log"
LOG_FILE.parent.mkdir(parents=True, exist_ok=True)

logging.basicConfig(
    filename=str(LOG_FILE),
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s"
)

app = Flask(__name__)
sock = Sock(app)

store = MemoryStore()
auth = AuthManager(store=store)

CONNECTED_CLIENTS = {}

UPLOADS_DIR.mkdir(parents=True, exist_ok=True)

ALLOWED_UPLOAD_EXTENSIONS = {
    ".txt", ".md", ".rtf", ".pdf", ".doc", ".docx", ".odt",
    ".csv", ".tsv", ".xls", ".xlsx", ".ods", ".json",
    ".ppt", ".pptx", ".odp",
    ".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".svg",
    ".py", ".js", ".ts", ".html", ".css", ".xml", ".yaml", ".yml",
    ".dwg", ".dxf", ".zip", ".sh", ".bash", ".ksh", ".zsh", ".ps1", ".bat"
}
MAX_UPLOAD_SIZE_BYTES = 50 * 1024 * 1024
SUPPORTED_LANGUAGES = {
    "auto", "en", "es", "fr", "de", "it", "pt", "nl", "pl", "ru",
    "uk", "hi", "ur", "ar", "tr", "ja", "ko", "zh"
}
SUPPORTED_UPLOAD_SCOPES = {"personal", "global"}


def _json_error(message: str, status: int = 400):
    return jsonify({"success": False, "message": message}), status


def _session_token_from_request():
    token = request.cookies.get("jarvis_mark3_session")
    if token:
        return token
    auth_header = request.headers.get("Authorization", "")
    if auth_header.startswith("Bearer "):
        return auth_header[7:].strip()
    return ""


def _current_user():
    token = _session_token_from_request()
    if not token:
        return None
    return auth.validate_session(token)


def _is_admin(user: dict) -> bool:
    if not user:
        return False
    return (user.get("email") or "").strip().lower() in {e.lower() for e in ADMIN_EMAILS}


def _require_admin():
    user = _current_user()
    if not user:
        return None, _json_error("Not authenticated.", 401)
    if not _is_admin(user):
        return None, _json_error("Admin access required.", 403)
    return user, None


def _normalize_language(value: str) -> str:
    lang = (value or "").strip().lower()
    if not lang:
        return "auto"
    if lang not in SUPPORTED_LANGUAGES:
        return "auto"
    return lang


def _normalize_upload_scope(value: str, user: dict) -> str:
    scope = (value or "").strip().lower()
    if scope not in SUPPORTED_UPLOAD_SCOPES:
        scope = "personal"
    if scope == "global" and not _is_admin(user):
        scope = "personal"
    return scope


def _resolved_display_name(user_dict: dict):
    settings = user_dict.get("settings") or {}
    preferred = (settings.get("preferred_name") or "").strip()
    display_name = (user_dict.get("display_name") or "").strip()
    return preferred or display_name or DEFAULT_DISPLAY_NAME


def _user_payload(user_dict: dict):
    settings = dict(user_dict.get("settings") or {})
    resolved_name = _resolved_display_name(user_dict)
    settings["preferred_name"] = resolved_name
    settings["preferred_language"] = _normalize_language(settings.get("preferred_language"))
    settings.setdefault("first_name", "")
    settings.setdefault("last_name", "")
    settings.setdefault("home_location", "")
    settings.setdefault("home_location_label", settings.get("home_location", ""))
    settings.setdefault("home_location_lat", None)
    settings.setdefault("home_location_lon", None)
    return {
        "email": user_dict.get("email", ""),
        "display_name": resolved_name,
        "settings": settings,
    }


def _admin_user_payload(row):
    return {
        "id": row["id"],
        "email": row["email"],
        "display_name": row["display_name"],
        "preferred_name": row["preferred_name"],
        "first_name": row["first_name"],
        "last_name": row["last_name"],
        "preferred_language": _normalize_language(row["preferred_language"]),
        "home_location": row["home_location"],
        "home_location_label": row["home_location_label"],
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
        "last_login_at": row["last_login_at"],
        "is_active": row["is_active"],
    }


def _safe_user_upload_dir(user: dict, scope: str = "personal") -> Path:
    if scope == "global":
        target = UPLOADS_DIR / "_global"
    else:
        user_id = str(user.get("user_id") or "anonymous")
        target = UPLOADS_DIR / user_id
    target.mkdir(parents=True, exist_ok=True)
    return target


def _sanitize_upload_name(filename: str) -> str:
    cleaned = secure_filename(filename or "")
    if not cleaned:
        cleaned = f"upload_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.bin"
    return cleaned


def _is_allowed_upload(filename: str) -> bool:
    suffix = Path(filename or "").suffix.lower()
    return suffix in ALLOWED_UPLOAD_EXTENSIONS


def _unique_destination(directory: Path, filename: str) -> Path:
    candidate = directory / filename
    if not candidate.exists():
        return candidate

    stem = candidate.stem
    suffix = candidate.suffix
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    counter = 1
    while True:
        alt = directory / f"{stem}_{timestamp}_{counter}{suffix}"
        if not alt.exists():
            return alt
        counter += 1


def _extract_plain_text_file(file_path: Path) -> str:
    return file_path.read_text(encoding="utf-8", errors="ignore")[:200000]


def _extract_docx_via_python_docx(file_path: Path) -> str:
    try:
        from docx import Document
        doc = Document(str(file_path))
        lines = []

        for p in doc.paragraphs:
            text = (p.text or "").strip()
            if text:
                lines.append(text)

        for table in doc.tables:
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    cell_text = " ".join(
                        (para.text or "").strip()
                        for para in cell.paragraphs
                        if (para.text or "").strip()
                    ).strip()
                    if cell_text:
                        cells.append(cell_text)
                if cells:
                    lines.append(" | ".join(cells))

        return "\n".join(lines)[:200000]
    except Exception:
        logging.exception("python-docx extraction failed for %s", file_path)
        return ""


def _extract_docx_via_zip_xml(file_path: Path) -> str:
    try:
        with zipfile.ZipFile(file_path, "r") as zf:
            if "word/document.xml" not in zf.namelist():
                return ""
            xml_data = zf.read("word/document.xml")

        root = ET.fromstring(xml_data)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

        paragraphs = []
        for para in root.findall(".//w:p", ns):
            texts = []
            for node in para.findall(".//w:t", ns):
                if node.text:
                    texts.append(node.text)
            joined = "".join(texts).strip()
            if joined:
                paragraphs.append(joined)

        return "\n".join(paragraphs)[:200000]
    except Exception:
        logging.exception("zip/xml docx extraction failed for %s", file_path)
        return ""


def _extract_pdf_text(file_path: Path) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(file_path))
        pages = []
        for page in reader.pages:
            txt = page.extract_text() or ""
            if txt.strip():
                pages.append(txt.strip())
        return "\n\n".join(pages)[:200000]
    except Exception:
        logging.exception("pypdf extraction failed for %s", file_path)
        return ""


def _extract_text_for_memory(file_path: Path, mime_type: str) -> str:
    suffix = file_path.suffix.lower()

    try:
        if suffix in {
            ".txt", ".md", ".rtf", ".json", ".csv", ".tsv",
            ".py", ".js", ".ts", ".html", ".css", ".xml",
            ".yaml", ".yml", ".sh", ".bash", ".ksh", ".zsh",
            ".ps1", ".bat"
        }:
            return _extract_plain_text_file(file_path)
    except Exception:
        logging.exception("Failed reading plain text upload: %s", file_path)

    if suffix == ".docx":
        text = _extract_docx_via_python_docx(file_path)
        if text.strip():
            return text
        text = _extract_docx_via_zip_xml(file_path)
        if text.strip():
            return text

    if suffix == ".pdf":
        text = _extract_pdf_text(file_path)
        if text.strip():
            return text

    return ""


def _chunk_text(text: str, chunk_size: int = 1000, overlap: int = 150):
    clean = (text or "").strip()
    if not clean:
        return []

    chunks = []
    start = 0
    length = len(clean)
    while start < length:
        end = min(start + chunk_size, length)
        chunk = clean[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= length:
            break
        start = max(end - overlap, start + 1)
    return chunks


def _write_upload_sidecar_meta(user: dict, file_path: Path, original_name: str, mime_type: str, extracted_text: str, scope: str):
    meta_path = file_path.with_suffix(file_path.suffix + ".meta.json")
    payload = {
        "user_id": user.get("user_id"),
        "email": user.get("email"),
        "display_name": _resolved_display_name(user),
        "original_name": original_name,
        "stored_name": file_path.name,
        "stored_path": str(file_path),
        "mime_type": mime_type,
        "scope": scope,
        "size_bytes": file_path.stat().st_size if file_path.exists() else 0,
        "uploaded_at": datetime.utcnow().isoformat() + "Z",
        "text_extracted": bool(extracted_text.strip()),
        "text_preview": extracted_text[:2000] if extracted_text else "",
    }
    meta_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    return meta_path


def _register_with_memory_store(user: dict, original_name: str, file_path: Path, mime_type: str, extracted_text: str) -> bool:
    user_id = user.get("user_id")
    success = False

    try:
        store.store_upload(
            user_id=user_id,
            file_name=original_name,
            stored_path=str(file_path),
            mime_type=mime_type
        )
        success = True
    except Exception:
        logging.exception("store.store_upload failed")

    if extracted_text.strip():
        fact_text = (
            f"Uploaded document: {original_name}\n"
            f"Path: {file_path}\n"
            f"Document contents excerpt:\n{extracted_text[:6000]}"
        )
        try:
            store.add_long_term_fact(
                user_id=user_id,
                fact_text=fact_text,
                confidence=0.95
            )
            success = True
        except Exception:
            logging.exception("store.add_long_term_fact failed")
    else:
        fact_text = (
            f"User uploaded document '{original_name}' stored at '{file_path}'. "
            f"No text extraction was available for this file type."
        )
        try:
            store.add_long_term_fact(
                user_id=user_id,
                fact_text=fact_text,
                confidence=0.60
            )
            success = True
        except Exception:
            logging.exception("store.add_long_term_fact fallback failed")

    return success


def _index_upload_into_memory(user: dict, file_path: Path, original_name: str, mime_type: str, scope: str) -> dict:
    extracted_text = _extract_text_for_memory(file_path, mime_type)
    meta_path = _write_upload_sidecar_meta(user, file_path, original_name, mime_type, extracted_text, scope)
    chunks = _chunk_text(extracted_text)

    document_id = None
    try:
        document_id = store.create_document_index(
            owner_user_id=user.get("user_id"),
            file_name=original_name,
            stored_path=str(file_path),
            mime_type=mime_type,
            scope=scope,
            extracted_text=extracted_text,
            ocr_status="complete" if extracted_text.strip() else "limited",
            is_searchable=1,
        )
        store.replace_document_chunks(
            document_id=document_id,
            owner_user_id=user.get("user_id"),
            scope=scope,
            chunks=chunks
        )
    except Exception:
        logging.exception("Document index storage failed")

    registered = _register_with_memory_store(
        user=user,
        original_name=original_name,
        file_path=file_path,
        mime_type=mime_type,
        extracted_text=extracted_text
    )

    if extracted_text.strip():
        if registered:
            return {
                "summary": f"{original_name} uploaded and indexed for Jarvis questions.",
                "text_extracted": True,
                "text_preview": extracted_text[:500],
                "registered": True,
                "document_id": document_id,
                "scope": scope,
                "chunk_count": len(chunks),
            }
        return {
            "summary": f"{original_name} uploaded successfully. Text was extracted, but registration into memory was limited. Metadata was saved at {meta_path}.",
            "text_extracted": True,
            "text_preview": extracted_text[:500],
            "registered": False,
            "document_id": document_id,
            "scope": scope,
            "chunk_count": len(chunks),
        }

    if registered:
        return {
            "summary": f"{original_name} uploaded. File registered, but text extraction is limited for this format.",
            "text_extracted": False,
            "text_preview": "",
            "registered": True,
            "document_id": document_id,
            "scope": scope,
            "chunk_count": len(chunks),
        }

    return {
        "summary": f"{original_name} uploaded successfully. File was saved, but text extraction/registration is limited for this format. Metadata was saved at {meta_path}.",
        "text_extracted": False,
        "text_preview": "",
        "registered": False,
        "document_id": document_id,
        "scope": scope,
        "chunk_count": len(chunks),
    }


def socket_request(socket_path: Path, payload: dict, recv_size: int = 1024 * 1024) -> dict:
    with socket.socket(socket.AF_UNIX, socket.SOCK_STREAM) as client:
        client.connect(str(socket_path))
        client.sendall(json.dumps(payload).encode("utf-8"))
        raw = client.recv(recv_size)
        if not raw:
            return {"ok": False, "error": f"No response from {socket_path.name}."}
        return json.loads(raw.decode("utf-8", errors="ignore"))


def brain_request(payload: dict) -> dict:
    return socket_request(BRAIN_SOCKET, payload)


def ears_request(payload: dict) -> dict:
    return socket_request(EARS_SOCKET, payload)


def mouth_request(payload: dict) -> dict:
    return socket_request(MOUTH_SOCKET, payload, recv_size=1024 * 1024)


def eyes_request(payload: dict) -> dict:
    return socket_request(EYES_SOCKET, payload, recv_size=4 * 1024 * 1024)


def _vision_payload(user: dict, action: str) -> dict:
    payload = request.get_json(force=True, silent=True) or {}
    if not isinstance(payload, dict):
        payload = {}
    payload = dict(payload)
    payload["action"] = payload.get("action") or action
    payload["user_id"] = user.get("user_id")
    payload["email"] = user.get("email", "")
    payload["is_admin"] = _is_admin(user)
    return payload


def _vision_http_status(result: dict) -> int:
    try:
        status_code = int((result or {}).get("status_code") or 200)
    except (TypeError, ValueError, AttributeError):
        status_code = 200
    return status_code if 100 <= status_code <= 599 else 200


@app.route("/")
@app.route("/dev-mark3/")
def route_index():
    return send_from_directory(str(WEB_DIR), "index.html")


@app.route("/hud")
@app.route("/dev-mark3/hud")
def route_hud():
    return send_from_directory(str(WEB_DIR), "hud.html")


@app.route("/admin")
@app.route("/dev-mark3/admin")
def route_admin():
    user = _current_user()
    if not user:
        return send_from_directory(str(WEB_DIR), "index.html")
    if not _is_admin(user):
        return _json_error("Admin access required.", 403)
    return send_from_directory(str(WEB_DIR), "admin.html")


@app.route("/css/<path:filename>")
def route_css(filename):
    return send_from_directory(str(WEB_DIR / "css"), filename)


@app.route("/js/<path:filename>")
def route_js(filename):
    return send_from_directory(str(WEB_DIR / "js"), filename)


@app.route("/assets/<path:filename>")
def route_assets(filename):
    return send_from_directory(str(WEB_DIR / "assets"), filename)


@app.route("/bg_image")
def route_bg_image():
    return send_from_directory(str(WEB_DIR / "assets" / "bg"), "jarvis-web-bg2-animated.gif")


@app.route("/health")
@app.route("/dev-mark3/health")
def route_health():
    return jsonify({"status": "ok", "service": "jarvis-mark3-interface"})


@app.route("/api/me", methods=["GET"])
@app.route("/dev-mark3/api/me", methods=["GET"])
def route_me():
    user = _current_user()
    if not user:
        return _json_error("Not authenticated.", 401)
    payload = _user_payload(user)
    payload["is_admin"] = _is_admin(user)
    return jsonify({"success": True, "user": payload, "is_admin": _is_admin(user)})


@app.route("/api/admin/users", methods=["GET"])
@app.route("/dev-mark3/api/admin/users", methods=["GET"])
def route_admin_users():
    admin_user, error = _require_admin()
    if error:
        return error

    rows = store.list_users()
    return jsonify({
        "success": True,
        "users": [_admin_user_payload(row) for row in rows]
    })


@app.route("/api/admin/users/<int:user_id>", methods=["GET"])
@app.route("/dev-mark3/api/admin/users/<int:user_id>", methods=["GET"])
def route_admin_user_detail(user_id: int):
    admin_user, error = _require_admin()
    if error:
        return error

    row = store.get_user_admin_detail(user_id)
    if not row:
        return _json_error("User not found.", 404)

    return jsonify({
        "success": True,
        "user": {
            "id": row["id"],
            "email": row["email"],
            "display_name": row["display_name"],
            "preferred_name": row["preferred_name"],
            "first_name": row["first_name"],
            "last_name": row["last_name"],
            "preferred_language": _normalize_language(row["preferred_language"]),
            "home_location": row["home_location"],
            "home_location_label": row["home_location_label"],
            "home_location_lat": row["home_location_lat"],
            "home_location_lon": row["home_location_lon"],
            "created_at": row["created_at"],
            "updated_at": row["updated_at"],
            "last_login_at": row["last_login_at"],
            "is_active": row["is_active"],
            "settings": {
                "browser_audio_enabled": row["browser_audio_enabled"],
                "browser_mic_enabled": row["browser_mic_enabled"],
                "browser_mic_gain": row.get("browser_mic_gain"),
                "browser_noise_gate": row.get("browser_noise_gate"),
                "browser_speech_threshold": row.get("browser_speech_threshold"),
                "local_audio_enabled": row["local_audio_enabled"],
                "local_mic_enabled": row["local_mic_enabled"],
            }
        }
    })


@app.route("/api/admin/users/<int:user_id>/update", methods=["POST"])
@app.route("/dev-mark3/api/admin/users/<int:user_id>/update", methods=["POST"])
def route_admin_user_update(user_id: int):
    admin_user, error = _require_admin()
    if error:
        return error

    row = store.get_user_admin_detail(user_id)
    if not row:
        return _json_error("User not found.", 404)

    data = request.get_json(force=True, silent=True) or {}
    preferred_name = (data.get("preferred_name") or data.get("display_name") or "").strip()
    first_name = (data.get("first_name") or "").strip()
    last_name = (data.get("last_name") or "").strip()
    preferred_language = _normalize_language(data.get("preferred_language") or "auto")
    new_password = data.get("new_password") or ""
    confirm_password = data.get("confirm_password") or ""
    home_location = (data.get("home_location") or "").strip()
    activate = data.get("activate")
    deactivate = data.get("deactivate")

    if preferred_name:
        auth.update_display_name(user_id, preferred_name)

    if new_password or confirm_password:
        if not new_password or not confirm_password:
            return _json_error("Both password fields are required.", 400)
        if new_password != confirm_password:
            return _json_error("Passwords do not match.", 400)
        auth.update_password(user_id, new_password)

    settings_updates = {
        "preferred_name": preferred_name,
        "first_name": first_name,
        "last_name": last_name,
        "preferred_language": preferred_language,
    }

    if home_location or home_location == "":
        validated = validate_home_location(home_location)
        if not validated.get("ok"):
            return _json_error(validated.get("error", "Invalid home location."), 400)
        settings_updates.update(validated["normalized"])

    store.update_user_settings(user_id, **settings_updates)

    if activate:
        store.activate_user(user_id)

    if deactivate:
        if (row["email"] or "").strip().lower() == (admin_user["email"] or "").strip().lower():
            return _json_error("You cannot deactivate your own admin account.", 400)
        store.deactivate_user(user_id)

    fresh = store.get_user_admin_detail(user_id)
    return jsonify({
        "success": True,
        "message": "User updated.",
        "user": {
            "id": fresh["id"],
            "email": fresh["email"],
            "display_name": fresh["display_name"],
            "preferred_name": fresh["preferred_name"],
            "first_name": fresh["first_name"],
            "last_name": fresh["last_name"],
            "preferred_language": _normalize_language(fresh["preferred_language"]),
            "home_location": fresh["home_location"],
            "home_location_label": fresh["home_location_label"],
            "created_at": fresh["created_at"],
            "updated_at": fresh["updated_at"],
            "last_login_at": fresh["last_login_at"],
            "is_active": fresh["is_active"],
        }
    })


@app.route("/api/admin/users/<int:user_id>/delete", methods=["POST"])
@app.route("/dev-mark3/api/admin/users/<int:user_id>/delete", methods=["POST"])
def route_admin_user_delete(user_id: int):
    admin_user, error = _require_admin()
    if error:
        return error

    row = store.get_user_admin_detail(user_id)
    if not row:
        return _json_error("User not found.", 404)

    if (row["email"] or "").strip().lower() == (admin_user["email"] or "").strip().lower():
        return _json_error("You cannot delete your own admin account.", 400)

    store.delete_user(user_id)
    return jsonify({
        "success": True,
        "message": "User deleted."
    })


@app.route("/auth/register", methods=["POST"])
@app.route("/dev-mark3/auth/register", methods=["POST"])
def route_register():
    data = request.get_json(force=True, silent=True) or {}
    email = (data.get("email") or "").strip().lower()
    first_name = (data.get("first_name") or "").strip()
    last_name = (data.get("last_name") or "").strip()
    home_location = (data.get("home_location") or "").strip()
    preferred_language = _normalize_language(data.get("preferred_language") or "auto")
    password = data.get("password") or ""
    confirm_password = data.get("confirm_password") or ""
    preferred_name = (data.get("preferred_name") or "").strip()

    if not first_name:
        return _json_error("First name is required.", 400)
    if not last_name:
        return _json_error("Last name is required.", 400)
    if not password:
        return _json_error("Password is required.", 400)
    if password != confirm_password:
        return _json_error("Passwords do not match.", 400)

    display_name = preferred_name or first_name or DEFAULT_DISPLAY_NAME

    ok, message, user_id = auth.register_user(email, password, display_name)
    if not ok:
        return _json_error(message, 400)

    settings_updates = {
        "preferred_name": display_name,
        "first_name": first_name,
        "last_name": last_name,
        "preferred_language": preferred_language,
    }

    validated = validate_home_location(home_location)
    if home_location and not validated.get("ok"):
        return _json_error(validated.get("error", "Invalid home location."), 400)
    if validated.get("ok"):
        settings_updates.update(validated["normalized"])

    store.update_user_settings(user_id, **settings_updates)

    ok, _, payload = auth.login_user(email, password)
    if not ok or not payload:
        return _json_error("Registration succeeded but auto-login failed.", 500)

    response = make_response(jsonify({
        "success": True,
        "message": "Registration successful.",
        "session_token": payload["session_token"],
        "user": _user_payload(payload)
    }))
    response.set_cookie(
        "jarvis_mark3_session",
        payload["session_token"],
        max_age=60 * 60 * 24 * 7,
        httponly=False,
        samesite="Lax",
        secure=False,
        path="/"
    )
    return response


@app.route("/auth/login", methods=["POST"])
@app.route("/dev-mark3/auth/login", methods=["POST"])
def route_login():
    data = request.get_json(force=True, silent=True) or {}
    email = (data.get("email") or "").strip().lower()
    password = data.get("password") or ""

    ok, message, payload = auth.login_user(email, password)
    if not ok or not payload:
        return _json_error(message, 401)

    response = make_response(jsonify({
        "success": True,
        "message": message,
        "session_token": payload["session_token"],
        "user": _user_payload(payload)
    }))
    response.set_cookie(
        "jarvis_mark3_session",
        payload["session_token"],
        max_age=60 * 60 * 24 * 7,
        httponly=False,
        samesite="Lax",
        secure=False,
        path="/"
    )
    return response


@app.route("/auth/logout", methods=["POST"])
@app.route("/dev-mark3/auth/logout", methods=["POST"])
def route_logout():
    token = _session_token_from_request()
    if token:
        auth.logout_session(token)

    response = make_response(jsonify({"success": True, "message": "Logged out."}))
    response.delete_cookie("jarvis_mark3_session", path="/")
    return response


@app.route("/auth/recover", methods=["POST"])
@app.route("/dev-mark3/auth/recover", methods=["POST"])
def route_recover():
    data = request.get_json(force=True, silent=True) or {}
    email = (data.get("email") or "").strip().lower()

    ok, message, token = auth.create_reset_token(email)
    if not ok:
        return _json_error(message, 404)

    try:
        from mailer import dispatch_recovery_email
        email_sent = dispatch_recovery_email(email, token)
        if not email_sent:
            return _json_error("Reset token created, but email dispatch failed.", 500)
    except Exception:
        logging.exception("Failed dispatching recovery email")
        return _json_error("Reset token created, but email dispatch failed.", 500)

    response = {
        "success": True,
        "message": "Password reset email dispatched."
    }

    try:
        from config import DEV_MODE
        if DEV_MODE:
            response["debug_reset_token"] = token
    except Exception:
        pass

    return jsonify(response)


@app.route("/auth/reset", methods=["POST"])
@app.route("/dev-mark3/auth/reset", methods=["POST"])
def route_reset():
    data = request.get_json(force=True, silent=True) or {}
    token = (data.get("token") or "").strip()
    password = data.get("password") or ""

    if not token or not password:
        return _json_error("Reset token and password are required.", 400)

    ok, message = auth.use_reset_token(token, password)
    if not ok:
        return _json_error(message, 400)

    return jsonify({"success": True, "message": message})


@app.route("/auth/settings", methods=["POST"])
@app.route("/dev-mark3/auth/settings", methods=["POST"])
def route_settings():
    user = _current_user()
    if not user:
        return _json_error("Not authenticated.", 401)

    data = request.get_json(force=True, silent=True) or {}
    display_name = (data.get("display_name") or data.get("preferred_name") or "").strip()
    new_password = data.get("new_password") or ""
    confirm_password = data.get("confirm_password") or ""
    settings = dict(data.get("settings") or {})
    requested_home_location = (settings.get("home_location") or "").strip()
    settings["preferred_language"] = _normalize_language(settings.get("preferred_language") or "auto")

    if display_name:
        auth.update_display_name(user["user_id"], display_name)
        settings["preferred_name"] = display_name

    if new_password or confirm_password:
        if not new_password or not confirm_password:
            return _json_error("Both password fields are required to change your password.", 400)
        if new_password != confirm_password:
            return _json_error("Passwords do not match.", 400)
        auth.update_password(user["user_id"], new_password)

    if "home_location" in settings:
        validated = validate_home_location(requested_home_location)
        if not validated.get("ok"):
            return _json_error(validated.get("error", "Invalid home location."), 400)
        settings.update(validated["normalized"])

    allowed_settings = {
        "preferred_name": settings.get("preferred_name"),
        "first_name": settings.get("first_name"),
        "last_name": settings.get("last_name"),
        "preferred_language": settings.get("preferred_language"),
        "home_location": settings.get("home_location"),
        "home_location_label": settings.get("home_location_label"),
        "home_location_lat": settings.get("home_location_lat"),
        "home_location_lon": settings.get("home_location_lon"),
        "browser_audio_enabled": settings.get("browser_audio_enabled"),
        "browser_mic_enabled": settings.get("browser_mic_enabled"),
        "browser_mic_gain": settings.get("browser_mic_gain"),
        "browser_noise_gate": settings.get("browser_noise_gate"),
        "browser_speech_threshold": settings.get("browser_speech_threshold"),
        "local_audio_enabled": settings.get("local_audio_enabled"),
        "local_mic_enabled": settings.get("local_mic_enabled"),
    }
    clean_settings = {k: v for k, v in allowed_settings.items() if v is not None}
    if clean_settings:
        store.update_user_settings(user["user_id"], **clean_settings)

    fresh_user = auth.validate_session(user["session_token"])
    return jsonify({
        "success": True,
        "message": "Settings updated.",
        "user": _user_payload(fresh_user)
    })


@app.route("/api/upload", methods=["POST"])
@app.route("/dev-mark3/api/upload", methods=["POST"])
def route_upload():
    user = _current_user()
    if not user:
        return _json_error("Not authenticated.", 401)

    scope = _normalize_upload_scope(request.form.get("scope") or "personal", user)

    if "file" not in request.files:
        return _json_error("No file uploaded.", 400)

    incoming = request.files["file"]
    if not incoming or not incoming.filename:
        return _json_error("No file selected.", 400)

    original_name = incoming.filename.strip()
    safe_name = _sanitize_upload_name(original_name)

    if not _is_allowed_upload(safe_name):
        return _json_error("File type not allowed.", 400)

    content_length = request.content_length or 0
    if content_length > MAX_UPLOAD_SIZE_BYTES:
        return _json_error("File exceeds maximum size limit.", 413)

    user_dir = _safe_user_upload_dir(user, scope=scope)
    destination = _unique_destination(user_dir, safe_name)

    try:
        incoming.save(str(destination))
    except Exception:
        logging.exception("Failed saving uploaded file")
        return _json_error("Failed to save uploaded file.", 500)

    size_bytes = destination.stat().st_size if destination.exists() else 0
    mime_type = mimetypes.guess_type(str(destination))[0] or "application/octet-stream"

    try:
        indexing = _index_upload_into_memory(user, destination, original_name, mime_type, scope)
    except Exception:
        logging.exception("Failed indexing uploaded file")
        indexing = {
            "summary": f"{original_name} uploaded successfully, but indexing encountered an internal issue. Metadata was still saved.",
            "text_extracted": False,
            "text_preview": "",
            "registered": False,
            "document_id": None,
            "scope": scope,
            "chunk_count": 0,
        }

    logging.info(
        "User %s uploaded file %s -> %s (%s bytes, scope=%s)",
        user.get("email"),
        original_name,
        destination,
        size_bytes,
        scope
    )

    return jsonify({
        "success": True,
        "message": indexing["summary"],
        "summary": indexing["summary"],
        "file": {
            "original_name": original_name,
            "stored_name": destination.name,
            "stored_path": str(destination),
            "mime_type": mime_type,
            "size_bytes": size_bytes,
            "user_id": user.get("user_id"),
            "scope": scope,
            "document_id": indexing.get("document_id"),
            "chunk_count": indexing.get("chunk_count", 0),
        },
        "indexed": indexing.get("text_extracted", False),
        "registered": indexing.get("registered", False),
        "preview": indexing.get("text_preview", ""),
    })


@app.route("/api/uploads", methods=["GET"])
@app.route("/dev-mark3/api/uploads", methods=["GET"])
def route_list_uploads():
    user = _current_user()
    if not user:
        return _json_error("Not authenticated.", 401)

    docs = store.list_documents_for_user(user["user_id"])
    return jsonify({
        "success": True,
        "uploads": docs
    })


@app.route("/api/uploads/<path:filename>", methods=["GET"])
@app.route("/dev-mark3/api/uploads/<path:filename>", methods=["GET"])
def route_get_upload(filename):
    user = _current_user()
    if not user:
        return _json_error("Not authenticated.", 401)

    safe_name = Path(filename).name

    personal_dir = _safe_user_upload_dir(user, scope="personal")
    personal_target = personal_dir / safe_name
    if personal_target.exists() and personal_target.is_file():
        return send_from_directory(str(personal_dir), safe_name, as_attachment=True)

    if _is_admin(user):
        global_dir = _safe_user_upload_dir(user, scope="global")
        global_target = global_dir / safe_name
        if global_target.exists() and global_target.is_file():
            return send_from_directory(str(global_dir), safe_name, as_attachment=True)

    return _json_error("Upload not found.", 404)


@app.route("/api/documents/search", methods=["GET"])
@app.route("/dev-mark3/api/documents/search", methods=["GET"])
def route_search_documents():
    user = _current_user()
    if not user:
        return _json_error("Not authenticated.", 401)

    query = (request.args.get("q") or "").strip()
    if not query:
        return jsonify({"success": True, "results": []})

    results = store.search_document_chunks_for_user(user["user_id"], query, limit=12)
    return jsonify({
        "success": True,
        "query": query,
        "results": results
    })


@app.route("/api/vision/cameras", methods=["POST"])
@app.route("/dev-mark3/api/vision/cameras", methods=["POST"])
def route_vision_cameras():
    user = _current_user()
    if not user:
        return _json_error("Not authenticated.", 401)
    try:
        result = eyes_request(_vision_payload(user, "list_cameras"))
    except Exception:
        logging.exception("Eyes list_cameras request failed")
        return _json_error("Vision service unavailable.", 503)
    return jsonify(result), _vision_http_status(result)


@app.route("/api/vision/preview", methods=["POST"])
@app.route("/dev-mark3/api/vision/preview", methods=["POST"])
def route_vision_preview():
    user = _current_user()
    if not user:
        return _json_error("Not authenticated.", 401)
    try:
        result = eyes_request(_vision_payload(user, "preview_frame_heartbeat"))
    except Exception:
        logging.exception("Eyes preview request failed")
        return _json_error("Vision preview unavailable.", 503)
    return jsonify(result), _vision_http_status(result)


@app.route("/api/vision/inspect", methods=["POST"])
@app.route("/dev-mark3/api/vision/inspect", methods=["POST"])
def route_vision_inspect():
    user = _current_user()
    if not user:
        return _json_error("Not authenticated.", 401)
    try:
        result = eyes_request(_vision_payload(user, "inspect_object"))
    except Exception:
        logging.exception("Eyes inspect request failed")
        return _json_error("Vision inspection unavailable.", 503)
    return jsonify(result), _vision_http_status(result)


@sock.route("/jarvis-websocket")
def websocket_jarvis(ws):
    user_id = None

    try:
        session_token = ""
        try:
            session_token = (
                (request.args.get("token") if request.args else None)
                or request.cookies.get("jarvis_mark3_session")
                or request.cookies.get("jarvis_user_token")
                or ""
            )
        except Exception:
            session_token = ""

        try:
            user = auth.validate_session(session_token) if session_token else None
        except Exception:
            logging.exception("WebSocket authentication bootstrap failed")
            user = None

        if not user:
            return

        display_name = _resolved_display_name(user)
        user_id = user["user_id"]
        user_settings = user.get("settings") or {}
        home_location = (
            user_settings.get("home_location_label")
            or user_settings.get("home_location")
            or ""
        )
        preferred_language = _normalize_language(user_settings.get("preferred_language"))
        CONNECTED_CLIENTS[user_id] = ws

        try:
            try:
                ears_request({
                    "action": "reset_stream",
                    "stream_id": f"user-{user_id}",
                    "language": None if preferred_language == "auto" else preferred_language,
                })
            except Exception:
                logging.exception("Ears reset_stream failed during websocket open")

            startup_result = brain_request({
                "mode": "startup_greeting",
                "user_id": user_id,
                "display_name": display_name,
                "preferred_name": display_name,
                "preferred_language": preferred_language,
                "home_location": home_location,
                "source": "system",
                "text": "__startup_greeting__"
            })
            startup_text = (
                startup_result.get("response_text")
                or startup_result.get("text")
                or render_greeting(display_name)
            )
        except Exception:
            logging.exception("Startup greeting request failed")
            startup_text = render_greeting(display_name)

        try:
            ws.send(json.dumps({
                "type": "response",
                "text": startup_text
            }))
        except Exception:
            pass

        try:
            audio_result = mouth_request({
                "mode": "speak",
                "text": startup_text,
                "user_id": user_id
            })
            audio_bytes = audio_result.get("audio_bytes")
            if audio_bytes:
                ws.send(bytes(audio_bytes))
            else:
                pcm_file = audio_result.get("pcm_file")
                if pcm_file:
                    pcm_path = Path(pcm_file)
                    if pcm_path.exists():
                        ws.send(pcm_path.read_bytes())
        except Exception:
            logging.exception("Startup mouth request failed")

        while True:
            message = ws.receive()
            if message is None:
                logging.info("Websocket received None for user %s", user_id)
                break

            logging.info("Websocket message type for user %s: %s", user_id, type(message).__name__)

            if isinstance(message, bytes):
                logging.info("Received audio bytes from user %s: %s bytes", user_id, len(message))
                try:
                    transcript_payload = {
                        "action": "process_audio",
                        "stream_id": f"user-{user_id}",
                        "audio": message.hex(),
                        "wake_word": "jarvis",
                        "wake_required": True,
                    }
                    if preferred_language != "auto":
                        transcript_payload["language"] = preferred_language

                    transcript_result = ears_request(transcript_payload)
                    logging.info("Ears transcript result for user %s: %s", user_id, transcript_result)
                except Exception:
                    logging.exception("Ears request failed")
                    try:
                        ws.send(json.dumps({
                            "type": "error",
                            "message": "Speech processing failed."
                        }))
                    except Exception:
                        pass
                    continue

                transcript = (transcript_result.get("transcript") or "").strip()
                result_type = transcript_result.get("type") or ""
                is_final = result_type == "final" or bool(transcript_result.get("is_final", False))

                if not transcript:
                    continue

                if not is_final:
                    try:
                        ws.send(json.dumps({
                            "type": "partial_transcript",
                            "text": transcript,
                            "partial_transcript": transcript
                        }))
                    except Exception:
                        pass
                    continue

                wake_detected = bool(transcript_result.get("wake_word_detected", False))
                processed_text = (transcript_result.get("processed_text") or transcript).strip()

                if not wake_detected:
                    try:
                        ws.send(json.dumps({
                            "type": "status",
                            "message": "Wake word not detected."
                        }))
                    except Exception:
                        pass
                    continue

                if processed_text:
                    try:
                        ws.send(json.dumps({
                            "type": "transcript",
                            "text": processed_text,
                            "is_final": True
                        }))
                    except Exception:
                        pass
                else:
                    continue

                try:
                    brain_result = brain_request({
                        "mode": "chat",
                        "user_id": user_id,
                        "text": processed_text,
                        "display_name": display_name,
                        "preferred_name": display_name,
                        "preferred_language": preferred_language,
                        "home_location": home_location,
                        "source": "voice"
                    })
                except Exception:
                    logging.exception("Brain request failed")
                    try:
                        ws.send(json.dumps({
                            "type": "error",
                            "message": "Brain processing failed."
                        }))
                    except Exception:
                        pass
                    continue

                reply_text = (brain_result.get("response_text") or brain_result.get("reply") or brain_result.get("text") or "").strip()
                if reply_text:
                    try:
                        outbound = {
                            "type": "response",
                            "text": reply_text
                        }
                        if brain_result.get("action"):
                            outbound["action"] = brain_result.get("action")
                        if brain_result.get("document"):
                            outbound["document"] = brain_result.get("document")
                        ws.send(json.dumps(outbound))
                    except Exception:
                        pass

                    try:
                        audio_result = mouth_request({
                            "mode": "speak",
                            "text": reply_text,
                            "user_id": user_id
                        })
                        audio_bytes = audio_result.get("audio_bytes")
                        if audio_bytes:
                            ws.send(bytes(audio_bytes))
                        else:
                            pcm_file = audio_result.get("pcm_file")
                            if pcm_file:
                                pcm_path = Path(pcm_file)
                                if pcm_path.exists():
                                    ws.send(pcm_path.read_bytes())
                    except Exception:
                        logging.exception("Mouth request failed")
                continue

            try:
                payload = json.loads(message)
            except Exception:
                continue

            msg_type = payload.get("type")

            if msg_type == "text_override":
                text = (payload.get("text") or "").strip()
                if not text:
                    continue

                try:
                    brain_result = brain_request({
                        "mode": "chat",
                        "user_id": user_id,
                        "text": text,
                        "display_name": display_name,
                        "preferred_name": display_name,
                        "preferred_language": preferred_language,
                        "home_location": home_location,
                        "source": payload.get("source") or "manual_text"
                    })
                except Exception:
                    logging.exception("Brain request failed")
                    try:
                        ws.send(json.dumps({
                            "type": "error",
                            "message": "Brain processing failed."
                        }))
                    except Exception:
                        pass
                    continue

                reply_text = (brain_result.get("response_text") or brain_result.get("reply") or brain_result.get("text") or "").strip()
                if reply_text:
                    try:
                        outbound = {
                            "type": "response",
                            "text": reply_text
                        }
                        if brain_result.get("action"):
                            outbound["action"] = brain_result.get("action")
                        if brain_result.get("document"):
                            outbound["document"] = brain_result.get("document")
                        ws.send(json.dumps(outbound))
                    except Exception:
                        pass

                    try:
                        audio_result = mouth_request({
                            "mode": "speak",
                            "text": reply_text,
                            "user_id": user_id
                        })
                        audio_bytes = audio_result.get("audio_bytes")
                        if audio_bytes:
                            ws.send(bytes(audio_bytes))
                        else:
                            pcm_file = audio_result.get("pcm_file")
                            if pcm_file:
                                pcm_path = Path(pcm_file)
                                if pcm_path.exists():
                                    ws.send(pcm_path.read_bytes())
                    except Exception:
                        logging.exception("Mouth request failed")

            elif msg_type == "update_user_settings":
                display_name_in = (payload.get("display_name") or "").strip()
                preferred_language_in = _normalize_language(payload.get("preferred_language") or preferred_language)

                if display_name_in:
                    try:
                        auth.update_display_name(user_id, display_name_in)
                        display_name = display_name_in
                    except Exception:
                        logging.exception("Failed updating display name from websocket")

                preferred_language = preferred_language_in

                try:
                    store.update_user_settings(
                        user_id,
                        preferred_language=preferred_language
                    )
                except Exception:
                    logging.exception("Failed updating preferred language from websocket")

                try:
                    ws.send(json.dumps({
                        "type": "status",
                        "message": "Settings updated."
                    }))
                except Exception:
                    pass

    except Exception:
        logging.exception("Websocket error")
    finally:
        if user_id is not None and CONNECTED_CLIENTS.get(user_id) is ws:
            CONNECTED_CLIENTS.pop(user_id, None)


if __name__ == "__main__":
    app.run(host=HOST, port=PORT)
